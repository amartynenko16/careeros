"""Tests for cover_letter.py, using a synthetic template matching the real
one's structure (greeting, body paragraphs, "Best," sign-off)."""

from pathlib import Path

import docx
import pytest

from careeros import cover_letter, resume


@pytest.fixture
def template_path(tmp_path: Path) -> Path:
    path = tmp_path / "cover_letter_template.docx"
    doc = docx.Document()
    doc.add_paragraph("")
    doc.add_paragraph("Hi Old Company Team,")
    doc.add_paragraph("Old hook paragraph.")
    doc.add_paragraph("Old evidence paragraph one.")
    doc.add_paragraph("Old evidence paragraph two.")
    doc.add_paragraph("Old closing ask.")
    doc.add_paragraph("Best,")
    doc.add_paragraph("Sample Candidate")
    doc.save(str(path))
    return path


def test_set_greeting(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    cover_letter.set_greeting(doc, "Docker")
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    assert any(p.text == "Hi Docker Team," for p in doc2.paragraphs)


def test_set_body_paragraphs_same_count(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    new_body = ["New hook.", "New evidence one.", "New evidence two.", "New ask."]
    cover_letter.set_body_paragraphs(doc, new_body)
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    greeting_idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.startswith("Hi "))
    signoff_idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.strip() == "Best,")
    body = [doc2.paragraphs[i].text for i in range(greeting_idx + 1, signoff_idx)]
    assert body == new_body
    # sign-off untouched
    assert doc2.paragraphs[signoff_idx + 1].text == "Sample Candidate"


def test_set_body_paragraphs_fewer(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    new_body = ["Only one paragraph now."]
    cover_letter.set_body_paragraphs(doc, new_body)
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    greeting_idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.startswith("Hi "))
    signoff_idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.strip() == "Best,")
    body = [doc2.paragraphs[i].text for i in range(greeting_idx + 1, signoff_idx)]
    assert body == new_body


def test_set_body_paragraphs_more(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    new_body = ["P1.", "P2.", "P3.", "P4.", "P5.", "P6."]
    cover_letter.set_body_paragraphs(doc, new_body)
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    greeting_idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.startswith("Hi "))
    signoff_idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.strip() == "Best,")
    body = [doc2.paragraphs[i].text for i in range(greeting_idx + 1, signoff_idx)]
    assert body == new_body


def test_set_body_paragraphs_empty_raises(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    with pytest.raises(resume.ResumeBuildError):
        cover_letter.set_body_paragraphs(doc, [])


def test_set_greeting_missing_raises() -> None:
    doc = docx.Document()
    doc.add_paragraph("no greeting here")
    with pytest.raises(resume.ResumeBuildError):
        cover_letter.set_greeting(doc, "Docker")
