"""Tests for resume.py's DOCX-editing logic.

Builds a small synthetic .docx that reproduces the structural quirks found
in a real-world template (multi-run paragraphs, a duplicated Professional
Summary label simulating the AlternateContent shape fallback) rather than
depending on anyone's personal file, which lives outside the repo.
"""

from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from careeros import db, resume


def _isolate_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Point the DB at a temp file so tests don't touch the real careeros.db.
    default_filename() looks up the company's display name via db.connect()
    and the person's name via facts.get(), which need a real (if empty)
    schema to query against."""
    from careeros import config

    monkeypatch.setattr(config, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    db.init_db()


def _seed_name(name: str) -> None:
    """Insert personal.name into immutable_facts the way facts.load_from_yaml
    would, without needing a real facts.yaml on disk."""
    import json
    from datetime import datetime, timezone

    with db.connect() as conn:
        conn.execute(
            "INSERT INTO immutable_facts (key, value_json, updated_at) VALUES (?, ?, ?)",
            ("personal.name", json.dumps(name), datetime.now(timezone.utc).isoformat(timespec="seconds")),
        )
        conn.commit()


def _add_multirun_paragraph(doc: docx.Document, texts: list[str], style: str | None = None) -> None:
    """Add a paragraph whose visible text is split across multiple runs,
    reproducing the old-formatting-artifact quirk seen in the real template."""
    p = doc.add_paragraph(style=style)
    for t in texts:
        p.add_run(t)


def _build_synthetic_template(path: Path) -> None:
    doc = docx.Document()

    doc.add_paragraph("TOP ACHIEVEMENTS")
    doc.add_paragraph("Some achievement.", style="List Paragraph")

    doc.add_paragraph("HARD SKILLS")
    doc.add_paragraph("Solution Architecture & Integration", style="List Paragraph")
    doc.add_paragraph("REST APIs, SDKs, JSON", style="List Paragraph")

    doc.add_paragraph("PERSONAL PROJECTS")
    doc.add_paragraph("Career OS")
    doc.add_paragraph("A Python project.")
    doc.add_paragraph("")  # blank spacer ends the block
    # Reproduces a real quirk found in a real-world template (2026-09-07):
    # an empty "List Paragraph"-styled paragraph sitting right before
    # EXPERIENCE, which can render as a stray floating bullet with no text.
    doc.add_paragraph("", style="List Paragraph")

    doc.add_paragraph("EXPERIENCE")
    header = doc.add_paragraph()
    header.add_run("SENIOR SOLUTIONS CONSULTANT • Acme Corp")
    doc.add_paragraph("Old bullet one.", style="List Paragraph")
    _add_multirun_paragraph(doc, ["Old bullet two, ", "split ", "across runs."], style="List Paragraph")
    doc.add_paragraph("Old bullet three.", style="List Paragraph")

    header2 = doc.add_paragraph()
    header2.add_run("TECHNICAL ACCOUNT MANAGER • Globex Inc")
    doc.add_paragraph("Globex Inc bullet one.", style="List Paragraph")
    doc.add_paragraph("Globex Inc bullet two.", style="List Paragraph")

    doc.save(str(path))

    # Now inject a duplicated "PROFESSIONAL SUMMARY" text-box-like structure
    # directly into the XML, since python-docx's high-level API can't create
    # floating text boxes -- but resume.py only needs two plain <w:p>
    # elements bearing the label, each followed by a content paragraph, to
    # exercise the same code path as the real duplicated shape.
    doc2 = docx.Document(str(path))
    body = doc2.element.body
    from docx.oxml import OxmlElement

    def _make_summary_block(text: str):
        label_p = OxmlElement("w:p")
        label_r = OxmlElement("w:r")
        label_t = OxmlElement("w:t")
        label_t.text = "PROFESSIONAL SUMMARY"
        label_r.append(label_t)
        label_p.append(label_r)

        content_p = OxmlElement("w:p")
        content_r = OxmlElement("w:r")
        content_t = OxmlElement("w:t")
        content_t.text = text
        content_r.append(content_t)
        content_p.append(content_r)
        return label_p, content_p

    old_text = "Old summary text."
    for _ in range(2):  # simulate the real template's two occurrences
        label_p, content_p = _make_summary_block(old_text)
        body.insert(0, content_p)
        body.insert(0, label_p)

    doc2.save(str(path))


@pytest.fixture
def template_path(tmp_path: Path) -> Path:
    path = tmp_path / "template.docx"
    _build_synthetic_template(path)
    return path


def test_set_summary_replaces_all_duplicated_labels(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    resume.set_summary(doc, "New summary text.")
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    matches = [
        t.text for t in doc2.element.body.iter(qn("w:t"))
        if t.text and "summary text" in t.text
    ]
    assert matches == ["New summary text.", "New summary text."]


def test_set_summary_missing_label_raises(template_path: Path) -> None:
    doc = docx.Document()
    doc.add_paragraph("no summary label here")
    with pytest.raises(resume.ResumeBuildError):
        resume.set_summary(doc, "text")


def test_summary_line_count_matches_observed_wrap() -> None:
    # The three lines actually observed wrapping in the Linear resume
    # overflow (2026-09-07) before a calibration regression creeps back in.
    line1 = ("Senior post-sales technical leader with 7 years of B2B SaaS experience across TAM, "
             "Solutions Consulting, and Technical Integrations, owning enterprise customer")
    line2 = ("relationships from onboarding through renewal and expansion. Track record turning "
             "customer advocacy into shipped product changes, backed by hands-on")
    line3 = ("technical fluency across APIs, SSO, and integration architecture. Builds the "
             "onboarding and enablement playbooks that scale a Customer Success function across a")
    tail = "top-account portfolio."
    overflowing_summary = " ".join([line1, line2, line3, tail])
    assert resume._summary_line_count(overflowing_summary) > resume.SUMMARY_BOX_MAX_LINES


def test_set_summary_rejects_overflowing_text(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    too_long = "word " * 200  # far past 3 lines at 150 chars/line
    with pytest.raises(resume.ResumeBuildError):
        resume.set_summary(doc, too_long)


def test_set_summary_accepts_text_within_box(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    fits = "Short summary well within the box." * 3  # ~105 chars, comfortably under 3 lines
    resume.set_summary(doc, fits)  # should not raise


def test_strip_empty_list_paragraphs_removes_stray_bullet(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    before = [p for p in doc.paragraphs if p.style.name == "List Paragraph" and not p.text.strip()]
    assert before, "fixture should contain the stray empty bullet this test exercises"

    resume._strip_empty_list_paragraphs(doc)

    after = [p for p in doc.paragraphs if p.style.name == "List Paragraph" and not p.text.strip()]
    assert after == []
    # real bullets must survive
    assert any("Some achievement." in p.text for p in doc.paragraphs)
    assert any("Globex Inc bullet two." in p.text for p in doc.paragraphs)


def test_set_employer_bullets_shrinks(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    resume.set_employer_bullets(doc, "Acme Corp", ["New bullet one.", "New bullet two."])
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    texts = [p.text for p in doc2.paragraphs if p.text.strip() in ("New bullet one.", "New bullet two.")]
    assert texts == ["New bullet one.", "New bullet two."]
    # the old third bullet paragraph should be gone
    assert not any("Old bullet three" in p.text for p in doc2.paragraphs)
    # Globex Inc's untouched bullets survive
    assert any("Globex Inc bullet one." in p.text for p in doc2.paragraphs)
    assert any("Globex Inc bullet two." in p.text for p in doc2.paragraphs)


def test_set_employer_bullets_grows_and_clears_multirun_leftovers(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    new_bullets = ["Bullet A.", "Bullet B.", "Bullet C.", "Bullet D.", "Bullet E."]
    resume.set_employer_bullets(doc, "Acme Corp", new_bullets)
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    acme_idx = next(i for i, p in enumerate(doc2.paragraphs) if "Acme Corp" in p.text)
    bullets = []
    i = acme_idx + 1
    while i < len(doc2.paragraphs) and doc2.paragraphs[i].style.name == "List Paragraph":
        bullets.append(doc2.paragraphs[i].text)
        i += 1
    assert bullets == new_bullets
    # no leftover text from the old multi-run "Old bullet two" paragraph
    assert not any("split" in b or "across runs" in b for b in bullets)


def test_set_employer_bullets_unknown_employer_raises(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    with pytest.raises(resume.ResumeBuildError):
        resume.set_employer_bullets(doc, "Nonexistent Co", ["x"])


def test_set_personal_projects_grows(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    resume.set_personal_projects(doc, [
        {"title": "Career OS", "description": "Updated description."},
        {"title": "New Project", "description": "New description."},
    ])
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.strip() == "PERSONAL PROJECTS")
    block = []
    i = idx + 1
    while doc2.paragraphs[i].text.strip():
        block.append(doc2.paragraphs[i].text)
        i += 1
    assert block == ["Career OS", "Updated description.", "New Project", "New description."]


def test_set_hard_skills_replaces_existing_subcategory(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    resume.set_hard_skills_subcategory(doc, "Solution Architecture & Integration", "New content line")
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.strip() == "Solution Architecture & Integration")
    assert doc2.paragraphs[idx + 1].text == "New content line"


def test_set_hard_skills_inserts_new_subcategory(template_path: Path) -> None:
    doc = docx.Document(str(template_path))
    resume.set_hard_skills_subcategory(doc, "Identity & Enterprise Integration", "SSO/SAML, OAuth, SCIM")
    doc.save(str(template_path))

    doc2 = docx.Document(str(template_path))
    idx = next(i for i, p in enumerate(doc2.paragraphs) if p.text.strip() == "Identity & Enterprise Integration")
    assert doc2.paragraphs[idx + 1].text == "SSO/SAML, OAuth, SCIM"
    # original subcategory still present
    assert any(p.text.strip() == "Solution Architecture & Integration" for p in doc2.paragraphs)


def test_default_filename_uses_company_display_name(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Regression test: an earlier version of default_filename() raised
    NameError (missing `db` import) because no test exercised the DB-lookup
    path -- every other test in this file only touches a docx directly."""
    _isolate_db(tmp_path, monkeypatch)
    _seed_name("Jordan Rivera")
    with db.connect() as conn:
        conn.execute(
            "INSERT INTO companies (slug, name, tier, ats_type, ats_slug, notes, added_at) "
            "VALUES ('vantageanalytics', 'Vantage Analytics', 1, 'ashby', 'vantageanalytics', '', '2026-09-02')"
        )
        conn.commit()

    job = {"company_slug": "vantageanalytics", "title": "Solutions Architect"}
    assert resume.default_filename(job, "Resume") == "Jordan Rivera Resume - Solutions Architect - Vantage Analytics.docx"
    assert resume.default_filename(job, "Cover Letter") == "Jordan Rivera Cover Letter - Solutions Architect - Vantage Analytics.docx"


def test_default_filename_falls_back_to_titled_slug(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Unknown company slug (e.g. a manually-added job) still produces a
    usable filename instead of erroring."""
    _isolate_db(tmp_path, monkeypatch)
    _seed_name("Jordan Rivera")
    job = {"company_slug": "some-new-co", "title": "Technical Account Manager"}
    assert resume.default_filename(job, "Resume") == "Jordan Rivera Resume - Technical Account Manager - Some New Co.docx"


def test_default_filename_falls_back_to_candidate_without_facts(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """No facts.yaml loaded yet (fresh clone, before setup) still produces a
    usable filename instead of erroring."""
    _isolate_db(tmp_path, monkeypatch)
    job = {"company_slug": "some-new-co", "title": "Technical Account Manager"}
    assert resume.default_filename(job, "Resume") == "Candidate Resume - Technical Account Manager - Some New Co.docx"
